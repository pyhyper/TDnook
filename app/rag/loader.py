import os
import io
import fitz  # PyMuPDF
import logging
import pandas as pd
from docx import Document

logger = logging.getLogger(__name__)

_ocr_reader = None

def _get_ocr_model_dir():
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    model_dir = os.path.join(base_dir, "data", "easyocr_models")
    os.makedirs(model_dir, exist_ok=True)
    return model_dir

def get_ocr_reader():
    global _ocr_reader
    if _ocr_reader is None:
        try:
            import easyocr
            ocr_langs = os.getenv("OCR_LANGUAGES", "vi,en").split(",")
            ocr_langs = [lang.strip() for lang in ocr_langs if lang.strip()]
            model_dir = _get_ocr_model_dir()
            logger.info(f"Initializing EasyOCR reader with languages: {ocr_langs}, model_dir: {model_dir}")
            _ocr_reader = easyocr.Reader(
                ocr_langs,
                gpu=False,
                verbose=False,
                model_storage_directory=model_dir,
                user_network_directory=model_dir,
            )
            logger.info("EasyOCR reader initialized successfully.")
        except Exception as e:
            logger.error(f"Failed to initialize OCR reader: {e}", exc_info=True)
            raise
    return _ocr_reader

def extract_text_with_ocr(page) -> str:
    try:
        reader = get_ocr_reader()
        zoom = 2.0
        mat = fitz.Matrix(zoom, zoom)
        pix = page.get_pixmap(matrix=mat, alpha=False)
        img_bytes = pix.tobytes("png")

        results = reader.readtext(img_bytes, detail=0, paragraph=True)
        text = "\n".join([r for r in results if r and str(r).strip()])
        return text
    except Exception as e:
        logger.error(f"OCR failed on page {page.number + 1}: {e}", exc_info=True)
        return ""

def extract_pages_from_pdf(file_path: str) -> list[dict]:
    pages = []
    try:
        doc = fitz.open(file_path)
        filename = os.path.basename(file_path)

        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text().strip()
            source = "native"

            if not text:
                text = extract_text_with_ocr(page)
                source = "ocr"

            if text:
                pages.append({
                    "page_num": page_num + 1,
                    "text": text,
                    "filename": filename,
                    "source": source,
                })

        doc.close()
        return pages
    except Exception as e:
        logger.error(f"Error extracting pages from PDF {file_path}: {e}", exc_info=True)
        raise e

def extract_from_image(file_path: str) -> list[dict]:
    filename = os.path.basename(file_path)
    try:
        reader = get_ocr_reader()
        results = reader.readtext(file_path, detail=0, paragraph=True)
        text = "\n".join([r for r in results if r and str(r).strip()])
        if not text:
            return []
        return [{
            "page_num": 1,
            "text": text,
            "filename": filename,
            "source": "image_ocr"
        }]
    except Exception as e:
        logger.error(f"Error OCR image {file_path}: {e}", exc_info=True)
        raise e

def extract_from_csv(file_path: str) -> list[dict]:
    filename = os.path.basename(file_path)
    try:
        df = pd.read_csv(file_path)
        pages = []
        total_rows = len(df)
        chunk_row_size = 50
        
        for i in range(0, total_rows, chunk_row_size):
            chunk_df = df.iloc[i : i + chunk_row_size]
            table_text = chunk_df.to_markdown(index=False)
            pages.append({
                "page_num": (i // chunk_row_size) + 1,
                "text": f"Table: {filename} (Rows {i+1} to {min(i+chunk_row_size, total_rows)} of {total_rows})\n\n{table_text}",
                "filename": filename,
                "source": "csv"
            })
            
        if not pages:
            pages.append({
                "page_num": 1,
                "text": df.to_string(index=False),
                "filename": filename,
                "source": "csv"
            })
        return pages
    except Exception as e:
        logger.error(f"Error extracting CSV {file_path}: {e}", exc_info=True)
        raise e

def extract_from_excel(file_path: str) -> list[dict]:
    filename = os.path.basename(file_path)
    try:
        excel_file = pd.ExcelFile(file_path)
        pages = []
        page_idx = 1
        
        for sheet_name in excel_file.sheet_names:
            df = excel_file.parse(sheet_name)
            if df.empty:
                continue
            total_rows = len(df)
            chunk_row_size = 50
            for i in range(0, total_rows, chunk_row_size):
                chunk_df = df.iloc[i : i + chunk_row_size]
                table_text = chunk_df.to_markdown(index=False)
                pages.append({
                    "page_num": page_idx,
                    "text": f"Sheet: '{sheet_name}' in {filename} (Rows {i+1} to {min(i+chunk_row_size, total_rows)})\n\n{table_text}",
                    "filename": filename,
                    "source": f"excel:{sheet_name}"
                })
                page_idx += 1
        return pages
    except Exception as e:
        logger.error(f"Error extracting Excel {file_path}: {e}", exc_info=True)
        raise e

def extract_from_docx(file_path: str) -> list[dict]:
    filename = os.path.basename(file_path)
    try:
        doc = Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        table_texts = []
        for table in doc.tables:
            data = []
            for row in table.rows:
                data.append([cell.text.strip() for cell in row.cells])
            if data:
                tdf = pd.DataFrame(data[1:], columns=data[0] if len(data) > 1 else None)
                table_texts.append(tdf.to_markdown(index=False))
                
        full_text = "\n\n".join(paragraphs + table_texts)
        if not full_text:
            return []
            
        return [{
            "page_num": 1,
            "text": full_text,
            "filename": filename,
            "source": "docx"
        }]
    except Exception as e:
        logger.error(f"Error extracting docx {file_path}: {e}", exc_info=True)
        raise e

def extract_from_text(file_path: str) -> list[dict]:
    filename = os.path.basename(file_path)
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            text = f.read().strip()
        if not text:
            return []
        return [{
            "page_num": 1,
            "text": text,
            "filename": filename,
            "source": "text"
        }]
    except Exception as e:
        logger.error(f"Error extracting text file {file_path}: {e}", exc_info=True)
        raise e

def extract_pages_from_file(file_path: str) -> list[dict]:
    """
    Unified dispatcher to extract text/pages from any supported file type:
    PDF, Images (.png, .jpg, .jpeg), Excel (.xlsx, .xls), CSV (.csv), Word (.docx), TXT.
    """
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == ".pdf":
        return extract_pages_from_pdf(file_path)
    elif ext in [".png", ".jpg", ".jpeg", ".bmp", ".webp"]:
        return extract_from_image(file_path)
    elif ext in [".xlsx", ".xls"]:
        return extract_from_excel(file_path)
    elif ext == ".csv":
        return extract_from_csv(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_from_docx(file_path)
    elif ext in [".txt", ".md", ".json", ".log"]:
        return extract_from_text(file_path)
    else:
        return extract_from_text(file_path)
